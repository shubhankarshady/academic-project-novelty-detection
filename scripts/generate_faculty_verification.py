"""
Generate Domain-Wise Faculty Verification Excel File
=====================================================
Creates an Excel file with one sheet per domain.
Each sheet shows project pairs with titles, abstracts,
and key details side-by-side so faculty can verify similarity labels
without needing to look up project IDs.

Includes:
- Summary sheet with similarity breakdown
- Color-coded similarity ratings
- Common keyword overlap detection
- Faculty verdict, comments, and reviewer tracking columns
"""

import os
import shutil
import tempfile
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def generate_verification_excel():
    root_dir = Path(__file__).resolve().parent.parent
    src = root_dir / "data" / "raw" / "Academic_Project_Similarity_Synthetic_Dataset.xlsx"
    output_path = root_dir / "data" / "raw" / "Faculty_Verification_DomainWise.xlsx"

    # Use a temp file copy in case the original is opened in an editor
    tmp = os.path.join(tempfile.gettempdir(), "temp_dataset.xlsx")
    shutil.copy2(src, tmp)

    xl = pd.ExcelFile(tmp)
    proj = pd.read_excel(xl, "Projects")
    pairs = pd.read_excel(xl, "SimilarityPairs")
    xl.close()
    if os.path.exists(tmp):
        os.remove(tmp)

    # Project A details
    proj_a = proj[["project_id", "domain", "project_title", "abstract",
                   "keywords", "algorithm_model", "technology_tools", "dataset"]].copy()
    proj_a.columns = ["project_A_id", "domain_A", "title_A", "abstract_A",
                       "keywords_A", "algorithm_A", "tools_A", "dataset_A"]

    # Project B details
    proj_b = proj[["project_id", "domain", "project_title", "abstract",
                   "keywords", "algorithm_model", "technology_tools", "dataset"]].copy()
    proj_b.columns = ["project_B_id", "domain_B", "title_B", "abstract_B",
                       "keywords_B", "algorithm_B", "tools_B", "dataset_B"]

    merged = pairs.merge(proj_a, on="project_A_id").merge(proj_b, on="project_B_id")
    merged["is_same_domain"] = merged["domain_A"] == merged["domain_B"]

    def find_common_keywords(kw_a, kw_b):
        if pd.isna(kw_a) or pd.isna(kw_b):
            return ""
        set_a = set(str(kw_a).lower().replace(",", " ").split())
        set_b = set(str(kw_b).lower().replace(",", " ").split())
        common = set_a & set_b
        common = {w for w in common if len(w) > 2}
        return ", ".join(sorted(common)[:10]) if common else "None"

    merged["common_keywords"] = merged.apply(
        lambda r: find_common_keywords(r["keywords_A"], r["keywords_B"]), axis=1
    )

    output_cols = [
        "pair_id",
        "similarity_category",
        "expert_similarity_label",
        "project_A_id",
        "domain_A",
        "title_A",
        "abstract_A",
        "algorithm_A",
        "tools_A",
        "dataset_A",
        "project_B_id",
        "domain_B",
        "title_B",
        "abstract_B",
        "algorithm_B",
        "tools_B",
        "dataset_B",
        "common_keywords",
        "is_same_domain",
        "annotation_note",
    ]

    col_rename = {
        "pair_id": "Pair ID",
        "similarity_category": "Similarity Level",
        "expert_similarity_label": "Score (0-4)",
        "project_A_id": "Project A ID",
        "domain_A": "Domain A",
        "title_A": "Project A - Title",
        "abstract_A": "Project A - Abstract",
        "algorithm_A": "Project A - Algorithm",
        "tools_A": "Project A - Tech Stack",
        "dataset_A": "Project A - Dataset",
        "project_B_id": "Project B ID",
        "domain_B": "Domain B",
        "title_B": "Project B - Title",
        "abstract_B": "Project B - Abstract",
        "algorithm_B": "Project B - Algorithm",
        "tools_B": "Project B - Tech Stack",
        "dataset_B": "Project B - Dataset",
        "common_keywords": "Common Keywords",
        "is_same_domain": "Same Domain?",
        "annotation_note": "Annotation Note",
    }

    all_domains = sorted(proj["domain"].unique())

    def prepare_sheet_df(df):
        out = df[output_cols].copy()
        out = out.rename(columns=col_rename)
        out = out.sort_values("Score (0-4)", ascending=False).reset_index(drop=True)
        out.index += 1
        out.index.name = "Sr. No."
        out["Faculty Verdict (0-4)"] = ""
        out["Faculty Comments"] = ""
        out["Verified By"] = ""
        return out

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_data = []
        for domain in all_domains:
            domain_pairs = merged[
                (merged["domain_A"] == domain) | (merged["domain_B"] == domain)
            ]
            same_domain = domain_pairs[domain_pairs["is_same_domain"]]
            cross_domain = domain_pairs[~domain_pairs["is_same_domain"]]
            summary_data.append({
                "Domain": domain,
                "Total Pairs": len(domain_pairs),
                "Same-Domain Pairs": len(same_domain),
                "Cross-Domain Pairs": len(cross_domain),
                "Not Similar (0)": len(domain_pairs[domain_pairs["expert_similarity_label"] == 0]),
                "Slightly Similar (1)": len(domain_pairs[domain_pairs["expert_similarity_label"] == 1]),
                "Moderately Similar (2)": len(domain_pairs[domain_pairs["expert_similarity_label"] == 2]),
                "Highly Similar (3)": len(domain_pairs[domain_pairs["expert_similarity_label"] == 3]),
                "Nearly Duplicate (4)": len(domain_pairs[domain_pairs["expert_similarity_label"] == 4]),
            })
        summary_df = pd.DataFrame(summary_data)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)

        for domain in all_domains:
            domain_mask = (merged["domain_A"] == domain) | (merged["domain_B"] == domain)
            domain_pairs = merged[domain_mask].copy()
            if len(domain_pairs) == 0:
                continue
            sheet_df = prepare_sheet_df(domain_pairs)
            sheet_name = domain[:31]
            sheet_df.to_excel(writer, sheet_name=sheet_name)

        cross = merged[~merged["is_same_domain"]].copy()
        cross_df = prepare_sheet_df(cross)
        cross_df.to_excel(writer, sheet_name="Cross-Domain Pairs")

    wb = load_workbook(output_path)
    sim_colors = {
        "Not Similar": "92D050",
        "Slightly Similar": "FFFF00",
        "Moderately Similar": "FFC000",
        "Highly Similar": "FF6600",
        "Nearly Duplicate": "FF0000",
    }
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = thin_border

        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            header_val = ws.cell(1, col_idx).value or ""
            if "Abstract" in str(header_val):
                ws.column_dimensions[col_letter].width = 50
            elif "Title" in str(header_val):
                ws.column_dimensions[col_letter].width = 35
            elif "Algorithm" in str(header_val) or "Tech Stack" in str(header_val) or "Dataset" in str(header_val):
                ws.column_dimensions[col_letter].width = 30
            elif "Domain" in str(header_val):
                ws.column_dimensions[col_letter].width = 25
            elif "Comment" in str(header_val) or "Verdict" in str(header_val):
                ws.column_dimensions[col_letter].width = 20
            elif "Common Keywords" in str(header_val):
                ws.column_dimensions[col_letter].width = 25
            else:
                ws.column_dimensions[col_letter].width = 15

        if ws.title == "Summary":
            continue

        sim_col_idx = None
        for col_idx in range(1, ws.max_column + 1):
            if ws.cell(1, col_idx).value == "Similarity Level":
                sim_col_idx = col_idx
                break

        if sim_col_idx:
            for row_idx in range(2, ws.max_row + 1):
                sim_val = ws.cell(row_idx, sim_col_idx).value
                if sim_val in sim_colors:
                    fill = PatternFill(start_color=sim_colors[sim_val],
                                       end_color=sim_colors[sim_val], fill_type="solid")
                    ws.cell(row_idx, sim_col_idx).fill = fill
                    ws.cell(row_idx, sim_col_idx + 1).fill = fill

                for col_idx in range(1, ws.max_column + 1):
                    ws.cell(row_idx, col_idx).alignment = Alignment(wrap_text=True, vertical="top")
                    ws.cell(row_idx, col_idx).border = thin_border

        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions

    ws_sum = wb["Summary"]
    ws_sum.freeze_panes = "B2"
    ws_sum.auto_filter.ref = ws_sum.dimensions
    for col_idx in range(1, ws_sum.max_column + 1):
        col_letter = get_column_letter(col_idx)
        ws_sum.column_dimensions[col_letter].width = 40 if col_idx == 1 else 18

    wb.save(output_path)
    print(f"Generated: {output_path}")


if __name__ == "__main__":
    generate_verification_excel()
